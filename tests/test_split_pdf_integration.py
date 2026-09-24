"""Results/Excel/PDF contract exercised through the existing real Split CSV pipeline."""

from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from pypdf import PdfReader

from core.split_comparison import build_split_comparison_pair, calculate_complete_split_pair
from core.split_corrections import apply_split_pair_correction, weather_sync_ambient_conditions
from core.split_deviation_analysis import analyze_split_selected_deviations
from core.split_display import format_split_pair_label
from core.split_energy import calculate_split_energy
from core.split_results import consolidate_split_final_results
from core.split_vehicle_mass import normalize_split_vehicle_mass_data
from core.split_weather_context import synchronize_weather_for_split_runs
from data.loaders import carregar_dados_csv_robusto
from data.split_exporters import export_split_final_results_to_excel
from data.split_parser import default_split_interval_config, parse_split_sources
from data.weather_loader import read_weather_file
from pages.page_split_results import _directional_pair_values, _pair_rows, _render_summary
from reports.split_pdf_report import _measured_runs, export_split_final_results_to_pdf
from translations import get_translator
from utils.split_graphs import build_split_selected_plot_series
from utils.split_pdf_workflow import pdf_metadata, prefill_report_metadata


def real_pipeline_inputs(selected_count=None):
    """Use the Eliezer CSVs/meteo and the 1545 kg reference used by import tests."""
    root = Path(__file__).resolve().parents[1] / "sample_data" / "Split"
    config = default_split_interval_config()
    sources = []
    for interval in ("high", "low"):
        path = root / "coastdown" / f"split eliezer {interval}.csv"
        _, runs, _ = carregar_dados_csv_robusto(
            str(path), using_split_method=True, is_alta=interval == "high",
        )
        sources.append({"filename": path.name, "role": interval, "all_run_data": runs})
    parsed = parse_split_sources(sources, config)
    weather = read_weather_file(root / "meteo" / "AGRICULTR_SPLIT.csv")
    parsed, _ = synchronize_weather_for_split_runs(parsed, weather)
    vehicle = normalize_split_vehicle_mass_data({"effective_mass": 1545.0})
    pairs = []
    run_pairs = ((1, 2), (3, 4), (5, 6)) if selected_count is None else [
        (run, run + 1) for run in range(1, 2 * selected_count, 2)
    ]
    for plus, minus in run_pairs:
        records = {
            f"{interval}_{suffix}": next(r for r in parsed[interval] if r["run_id"] == run_id)
            for interval in ("high", "low") for suffix, run_id in (("plus", plus), ("minus", minus))
        }
        raw = calculate_complete_split_pair(**records, effective_mass=vehicle["effective_mass_kg"], config=config)
        corrected = apply_split_pair_correction(raw, weather_sync_ambient_conditions({
            key: record["weather_sync"] for key, record in records.items()
        }))
        pair = build_split_comparison_pair(corrected, pair_id=f"eliezer-{plus}-{minus}")
        # Keep a nonselected real pair in the source and reverse the selected order.
        pair["selected"] = selected_count is not None or plus != 3
        pairs.insert(0, pair)
    summary = consolidate_split_final_results(pairs)
    analysis = analyze_split_selected_deviations(summary["selected_pairs"])
    state = {"split_interval_config": config, "split_input_sources": sources,
             "vehicle_info": {**vehicle, "test_date": parsed["high"][0]["start_timestamp"].date()}}
    inputs = dict(
        final_results=summary, vehicle_data=state["vehicle_info"], deviation_analysis=analysis,
        graph_series=build_split_selected_plot_series(summary["selected_pairs"], sources),
        test_name="Split Eliezer | CSV pipeline", generated_at=datetime.now().astimezone(),
        test_metadata=pdf_metadata(prefill_report_metadata({}, state), get_translator("pt")),
    )
    return inputs, state, pairs


class SplitPdfIntegrationTests(unittest.TestCase):
    def test_real_pipeline_results_excel_pdf_agree(self):
        inputs, _, pairs = real_pipeline_inputs()
        before = deepcopy(inputs)
        summary = inputs["final_results"]
        selected = summary["selected_pairs"]
        self.assertEqual([p["id"] for p in selected], [pairs[0]["id"], pairs[2]["id"]])
        ui_rows = _pair_rows(selected, get_translator("pt"))
        self.assertEqual([row["Par"] for row in ui_rows], [format_split_pair_label(p) for p in selected])
        ui = Mock()
        ui.container.return_value = ui
        with patch("pages.page_split_results.st", ui):
            _render_summary(summary, inputs["deviation_analysis"]["time_summary"], get_translator("pt"))
        values = [call.args[1] for call in ui.metric.call_args_list]
        self.assertEqual(values, [str(len(selected)), f"{summary['mean_f0']:.4f}",
                                  f"{summary['mean_f2']:.6f}", f"{summary['mean_energy']:.4f}"])
        workbook = load_workbook(BytesIO(export_split_final_results_to_excel(
            final_results=summary, selected_pairs=selected, vehicle_data=inputs["vehicle_data"],
            deviation_analysis=inputs["deviation_analysis"],
        )), data_only=True)
        excel_summary = {row[0]: row[1] for row in workbook["Resumo Final"].iter_rows(values_only=True)}
        self.assertEqual(excel_summary["Quantidade de pares selecionados"], summary["num_pairs"])
        for label, key in (("F0 final [N]", "mean_f0"), ("F2 final [N/(km/h)²]", "mean_f2"),
                           ("Energia média [MJ/km]", "mean_energy")):
            self.assertAlmostEqual(excel_summary[label], summary[key])
        # The renderer must not import or call any application engineering path.
        with patch("core.split_corrections.calculate_split_energy", side_effect=AssertionError("PDF calculation")):
            pdf = PdfReader(BytesIO(export_split_final_results_to_pdf(**inputs)))
        self.assertEqual(len(pdf.pages), 3)
        first, measured, last = [page.extract_text() for page in pdf.pages]
        self.assertIn(f"\n{summary['num_pairs']}\n", first)
        self.assertNotIn("indisponíveis", measured)
        self.assertNotIn(pairs[1]["id"], last)
        self.assertLess(last.index("Par 1"), last.index("Par 2"))
        for key, precision in (("mean_f0", 4), ("mean_f2", 6), ("mean_energy", 4)):
            self.assertIn(f"{summary[key]:.{precision}f}", first)
            self.assertIn(f"{summary[key]:.{3 if key == 'mean_f0' else precision}f}", last.split("Resultado final", 1)[1])
        deviations = list(workbook["Análise de Desvios e Tempos"].iter_rows(values_only=True))
        for pair, row in zip(selected, ui_rows):
            excel_row = next(r for r in deviations if r[0] == format_split_pair_label(pair))
            for key, label, precision, column in (("F0_mean", "F0 (N)", 4, 1),
                    ("F2_mean", "F2 (N/(km/h)²)", 6, 3), ("energy", "Energia (MJ/km)", 4, 5)):
                self.assertEqual(row[label], f"{pair[key]:.{precision}f}")
                self.assertAlmostEqual(excel_row[column], pair[key])
                self.assertIn(f"{pair[key]:.{3 if key == 'F0_mean' else precision}f}", last)
            for suffix in ("plus", "minus"):
                expected = calculate_split_energy(pair[f"F0_{suffix}"], pair[f"F2_{suffix}"])["energy"]
                self.assertEqual(pair[f"energy_{suffix}"], expected)
                self.assertIn(f"{expected:.4f}", _directional_pair_values(pair, suffix).values())
                self.assertNotIn("Energia bruta", last)  # Raw energy is omitted entirely.
        for interval, rows in _measured_runs(selected).items():
            self.assertEqual(len(rows), 4)
            for row in rows:
                self.assertTrue(row["weather"]["matched"])
                self.assertIn(f"{row['record']['delta_t_s']:.2f}", measured)
        self.assertEqual(inputs, before)

    def test_five_real_selected_pairs_render_in_canonical_order(self):
        inputs, _, _ = real_pipeline_inputs(5)
        reader = PdfReader(BytesIO(export_split_final_results_to_pdf(**inputs)))
        selected = inputs["final_results"]["selected_pairs"]
        self.assertEqual(len(selected), 5)
        self.assertEqual(len(inputs["graph_series"]), 20)
        text = "\n".join(page.extract_text() for page in reader.pages[2:])
        for key, label in (("cv_f0_prime", "CV F0"), ("cv_f2_prime", "CV F2")):
            self.assertIn(f"{label} [%]: {inputs['final_results'][key]:.2f}", text)
        positions = [text.index("Par " + str(index)) for index in range(1, 6)]
        self.assertEqual(positions, sorted(positions))
        for pair in selected:
            for suffix in ("plus", "minus", "mean"):
                self.assertIn(f"{pair[f'f0_prime_{suffix}']:.3f}", text)
                self.assertIn(f"{pair[f'f2_prime_{suffix}']:.6f}", text)
            for key, precision in (("F0_mean", 4), ("F2_mean", 6), ("energy", 4)):
                self.assertIn(f"{pair[key]:.{3 if key == 'F0_mean' else precision}f}", text)
        for key, precision in (("mean_f0", 4), ("mean_f2", 6), ("mean_energy", 4)):
            self.assertIn(f"{inputs['final_results'][key]:.{3 if key == 'mean_f0' else precision}f}", text.split("Resultado final")[1])

    def test_five_pair_cards_align_and_charts_share_pair_styles_in_both_languages(self):
        from html import escape
        from reportlab.graphics.charts.lineplots import LinePlot
        from reportlab.graphics.shapes import Drawing
        from reportlab.platypus import Paragraph
        from reports.report_styles import PAGE_MARGIN, PAGE_SIZE, build_report_styles
        from reports.split_pdf_report import _run_tables

        inputs, _, _ = real_pipeline_inputs(5)
        before = deepcopy(inputs)
        selected = inputs["final_results"]["selected_pairs"]
        styles = build_report_styles()
        for language in ("pt", "en"):
            def tr(pt, en):
                return pt if language == "pt" else en
            def paragraph(value, style="Table"):
                return Paragraph(escape(str(value)).replace("\n", "<br/>"), styles[style])
            # Series arrival order must not assign a different color or pair number.
            story = _run_tables(selected, list(reversed(inputs["graph_series"])),
                                PAGE_SIZE[0] - 2 * PAGE_MARGIN, paragraph, tr)
            cards = [cell[0] for cell in story[2]._cellvalues[0][::2]]
            self.assertEqual(len(cards), 5)
            self.assertTrue(all(card._rowHeights == cards[0]._rowHeights for card in cards))
            for card in cards:
                self.assertTrue(all(height <= 14 for height in card._rowHeights[2:4]))
                self.assertLessEqual(card._rowHeights[-1], 14)
            drawings = [item for item in story if isinstance(item, Drawing)]
            self.assertEqual(len(drawings), 2)
            charts = [next(item for item in drawing.contents if isinstance(item, LinePlot)) for drawing in drawings]
            for drawing, chart, interval in zip(drawings, charts, ("high", "low")):
                self.assertGreater(drawing.height, 150)
                self.assertGreater(chart.height, 90)
                self.assertEqual(len(chart.data), 10)
                for index, pair in enumerate(selected):
                    for offset, suffix in enumerate(("plus", "minus")):
                        curve = next(item for item in inputs["graph_series"]
                                     if item["record"] == pair[f"{interval}_{suffix}"])
                        self.assertEqual(chart.data[2 * index + offset], list(zip(curve["times_s"], curve["speeds_kmh"])))
                    self.assertEqual(chart.lines[2 * index].strokeColor, chart.lines[2 * index + 1].strokeColor)
                    self.assertIsNone(chart.lines[2 * index].strokeDashArray)
                    self.assertEqual(chart.lines[2 * index + 1].strokeDashArray, [4, 2])
            for index in range(10):
                self.assertEqual(charts[0].lines[index].strokeColor, charts[1].lines[index].strokeColor)
            self.assertEqual(len({charts[0].lines[index].strokeColor.hexval() for index in range(0, 10, 2)}), 5)
            reader = PdfReader(BytesIO(export_split_final_results_to_pdf(**{**inputs, "language": language})))
            measured = reader.pages[1].extract_text()
            for index in range(1, 6):
                self.assertIn(tr("Par ", "Pair ") + str(index), measured)
            for label in ("90-85", "85-80", "80-75", "75-70", "45-40", "40-35", "High+", "Low+", "High-", "Low-"):
                self.assertIn(label, measured)
            self.assertNotIn(tr("Condições climáticas", "Climate conditions"), measured)
            self.assertNotIn("[kPa]", measured)
        self.assertEqual(inputs, before)

    def test_graphs_follow_selection_and_preserve_sources(self):
        inputs, state, _ = real_pipeline_inputs()
        selected = inputs["final_results"]["selected_pairs"]
        curves = build_split_selected_plot_series(selected + selected, state["split_input_sources"])
        self.assertEqual(curves, inputs["graph_series"])
        self.assertEqual(len(curves), 8)
        for curve in curves:
            self.assertEqual(curve["data_mode"], "interval_curve")
            self.assertAlmostEqual(curve["times_s"][-1], curve["record"]["delta_t_s"])
            self.assertEqual(curve["speeds_kmh"][0], curve["record"]["start_kmh"])
        self.assertEqual(build_split_selected_plot_series([], state["split_input_sources"]), [])
        other = deepcopy(selected[0])
        other["high_plus"]["filename"] = "different-source.csv"
        self.assertEqual(len(build_split_selected_plot_series([selected[0], other], [])), 5)


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 2 and sys.argv[1] == "--sample":
        inputs, _, _ = real_pipeline_inputs(5)
        output = Path("output/pdf/split_report_five_pairs.pdf")
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(export_split_final_results_to_pdf(**inputs))
        print(output)
    else:
        unittest.main()
